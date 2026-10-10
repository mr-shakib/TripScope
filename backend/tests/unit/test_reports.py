"""Report center units: formatting, document building, and PDF/XLSX/CSV rendering (FR-09, spec §12.1).

A fake analytics service returns small deterministic results, including hostile zone names, so every template
can be built and rendered without services and the outputs can be checked cell by cell.
"""

from __future__ import annotations

import csv
import io
import uuid
from datetime import UTC, date, datetime, timedelta
from typing import Any

import openpyxl
import pytest
from pypdf import PdfReader

from tripscope.analytics.filters import AnalyticsFilters, TimeSeriesQuery
from tripscope.analytics.labels import FLAG_LABELS, QUARANTINE_LABELS
from tripscope.analytics.metrics import METRICS
from tripscope.analytics.service import Coverage
from tripscope.core.errors import ValidationFailedError
from tripscope.metadata.models import DatasetPeriod
from tripscope.pipeline.rules import FLAGS, QUARANTINE_REASONS
from tripscope.reports import formatting as f
from tripscope.reports.builder import ReportBuilder
from tripscope.reports.document import ChartBlock, ReportDocument, TableBlock
from tripscope.reports.render_csv import COLUMNS, render_report_csv
from tripscope.reports.render_pdf import render_pdf
from tripscope.reports.render_xlsx import render_report_xlsx, sheet_name, write_extract_xlsx
from tripscope.reports.templates import TEMPLATES

EVIL_ZONE = '=HYPERLINK("http://evil.example","x")'
MARKUP_ZONE = "<b>Bold & Co</b>"
RUN_ID = uuid.UUID("11111111-2222-3333-4444-555555555555")
MARCH = [date(2025, 3, 1) + timedelta(days=i) for i in range(31)]


def _trips(day: date) -> int:
    return 1000 + 10 * day.day + (300 if day.isoweekday() >= 6 else 0)


TOTAL = sum(_trips(d) for d in MARCH)


class FakeAnalytics:
    """Duck-typed stand-in for AnalyticsService with one published month (March 2025)."""

    def __init__(self, *, empty: bool = False) -> None:
        self.empty = empty

    def coverage(self, dataset_id: str) -> Coverage:
        period = DatasetPeriod(
            dataset_id=dataset_id,
            data_period=date(2025, 3, 1),
            run_id=RUN_ID,
            row_count=TOTAL,
            min_pickup_date=date(2025, 3, 1),
            max_pickup_date=date(2025, 3, 31),
            published_at=datetime(2025, 4, 2, 9, 30, tzinfo=UTC),
        )
        return Coverage(dataset_id, "Test Yellow Taxi", "yellow", "NYC TLC Trip Record Data (test)", [period])

    def zones(self) -> dict[int, dict[str, Any]]:
        return {
            1: {"zone": "Midtown Center", "borough": "Manhattan", "is_geographic": True},
            2: {"zone": EVIL_ZONE, "borough": "Queens", "is_geographic": True},
            3: {"zone": MARKUP_ZONE, "borough": "Manhattan", "is_geographic": True},
            264: {"zone": "Unknown", "borough": "Unknown", "is_geographic": False},
        }

    def _meta(self, table: str = "trips_hourly_agg") -> dict[str, Any]:
        return {"source_table": table}

    def _kpis(self, total: int) -> dict[str, Any]:
        values = {
            "total_trips": total,
            "avg_daily_trips": total / 31,
            "total_recorded_amount": total * 28.0,
            "avg_total_amount": 28.0,
            "avg_trip_distance": 3.25,
            "avg_trip_duration_minutes": 15.5,
        }
        excluded = {"total_recorded_amount": 120, "avg_total_amount": 120, "avg_trip_distance": 40}
        return {
            m: {"value": values[m], "unit": METRICS[m].unit, "excluded_rows": excluded.get(m, 0)}
            for m in METRICS
        }

    def overview(self, filters: AnalyticsFilters, *, compare: str = "none") -> dict[str, Any]:
        if self.empty:
            return {"kpis": None, "data_state": "empty", "meta": self._meta()}
        result: dict[str, Any] = {
            "kpis": self._kpis(TOTAL),
            "data_state": "ok",
            "result_range": {"first_date": MARCH[0], "last_date": MARCH[-1], "days_with_data": 31},
            "meta": self._meta(),
        }
        if compare == "previous":
            result["comparison"] = {
                "available": True,
                "start_date": date(2025, 1, 29),
                "end_date": date(2025, 2, 28),
                "days": 31,
                "kpis": self._kpis(TOTAL - 4000),
                "reason": None,
            }
        return result

    def time_series(self, query: TimeSeriesQuery) -> dict[str, Any]:
        points = [{"bucket": d, "value": _trips(d), "trips": _trips(d)} for d in MARCH]
        return {"points": points, "meta": self._meta()}

    def breakdown(self, filters: AnalyticsFilters, *, metric: str, dimension: str, limit: int = 300) -> Any:
        if dimension == "hour":
            groups = [
                {
                    "key": h,
                    "value": 20.0 + h if metric != "total_trips" else None,
                    "trips": 2000 + 100 * h,
                    "label": f"{h:02d}:00",
                }
                for h in range(24)
            ]
        elif dimension in ("pickup_zone", "dropoff_zone"):
            zones = self.zones()
            groups = [
                {
                    "key": z,
                    "value": None,
                    "trips": trips,
                    "label": zones[z]["zone"] if z != 264 else "Unmapped (264)",
                    "borough": zones[z]["borough"] if z != 264 else None,
                    "mapped": z != 264,
                }
                for z, trips in ((1, 9000), (2, 7000), (3, 5000), (264, 100))
            ][:limit]
        elif dimension == "payment_type":
            amounts = {"avg_total_amount": [29.5, 24.0], "avg_trip_distance": [3.3, 3.1]}.get(
                metric, [None, None]
            )
            groups = [
                {"key": 1, "value": amounts[0], "trips": 25000, "label": "Credit card"},
                {"key": 2, "value": amounts[1], "trips": 8000, "label": "Cash"},
            ]
        else:
            raise AssertionError(f"unexpected dimension {dimension}")
        return {"groups": groups, "meta": self._meta()}

    def hour_weekday_matrix(self, filters: AnalyticsFilters, *, metric: str) -> dict[str, Any]:
        cells = [
            {"weekday": d, "hour": h, "value": None, "trips": 100 * d + h}
            for d in range(1, 8)
            for h in range(24)
        ]
        return {"cells": cells, "meta": self._meta()}

    def flows(self, filters: AnalyticsFilters, *, metric: str, limit: int) -> dict[str, Any]:
        flows = [
            {
                "pickup_zone": 1,
                "dropoff_zone": 2,
                "value": None,
                "trips": 600,
                "pickup_label": "Midtown Center",
                "pickup_borough": "Manhattan",
                "dropoff_label": EVIL_ZONE,
                "dropoff_borough": "Queens",
                "same_zone": False,
            },
            {
                "pickup_zone": 3,
                "dropoff_zone": 3,
                "value": None,
                "trips": 400,
                "pickup_label": MARKUP_ZONE,
                "pickup_borough": "Manhattan",
                "dropoff_label": MARKUP_ZONE,
                "dropoff_borough": "Manhattan",
                "same_zone": True,
            },
        ]
        return {"flows": flows, "meta": self._meta("taxi_trips")}

    def distribution(self, filters: AnalyticsFilters, *, metric: str) -> dict[str, Any]:
        width, cap = (1.0, 50.0) if metric == "trip_distance" else (5.0, 200.0)
        buckets = [
            {"start": i * width, "end": (i + 1) * width, "trips": 1000 // (i + 1), "open_ended": False}
            for i in range(10)
        ] + [{"start": cap, "end": None, "trips": 7, "open_ended": True}]
        return {
            "buckets": buckets,
            "summary": {
                "counted_trips": sum(b["trips"] for b in buckets),
                "excluded_trips": 40,
                "bucket_width": width,
                "cap": cap,
                "median_bucket": {"start": width, "end": 2 * width},
                "p90_bucket": {"start": 6 * width, "end": 7 * width},
                "above_cap_trips": 7,
            },
            "meta": self._meta("fare_distance_buckets"),
        }

    def quality(self, filters: AnalyticsFilters) -> dict[str, Any]:
        period = {
            "period": "2025-03",
            "run_id": str(RUN_ID),
            "input_rows": TOTAL + 12,
            "accepted_rows": TOTAL,
            "quarantined_rows": 12,
            "duplicate_rows": 0,
            "quarantine_reasons": {"dropoff_before_pickup": 9, "pickup_outside_source_period": 3},
            "flags": {"passenger_count_missing_or_zero": 900, "invalid_amount": 120},
            "duration_seconds": 37.2,
            "schema_version": "yellow-2463cafe",
        }
        other = period | {"period": "2025-04", "input_rows": 99_999}  # outside the report: must be ignored
        daily = [
            {"date": d, "flag": flag, "flagged_trips": n, "trips": _trips(d)}
            for d in MARCH
            for flag, n in (("passenger_count_missing_or_zero", 30), ("invalid_amount", 4))
        ]
        return {"periods": [period, other], "totals": {}, "daily_flags": daily, "meta": self._meta()}

    def schema(self, dataset_id: str) -> dict[str, Any]:
        return {
            "schema_versions": [
                {
                    "schema_version": "yellow-2463cafe",
                    "file_format": "parquet",
                    "column_count": 20,
                    "periods": ["2025-03"],
                }
            ],
            "drift": [],
        }


def build(template: str, filters: AnalyticsFilters | None = None, **kw: Any) -> ReportDocument:
    spec = TEMPLATES[template]
    return ReportBuilder(FakeAnalytics(**kw)).build(  # type: ignore[arg-type]
        template=template,
        title=f"{spec.title} test",
        filters=filters or AnalyticsFilters(start_date=date(2025, 3, 1), end_date=date(2025, 3, 31)),
        sections=spec.section_ids(),
        prepared_by="Test Analyst",
        generated_at=datetime(2026, 10, 10, 12, 0, tzinfo=UTC),
    )


# ---- formatting ------------------------------------------------------------------------------------------


def test_display_matches_the_dashboard_conventions() -> None:
    assert f.display(4145143, "trips") == "4,145,143"
    assert f.display(110111232.044, "usd") == "$110,111,232.04"
    assert f.display(-5.5, "usd") == "-$5.50"
    assert f.display(3.344, "miles") == "3.34 mi"
    assert f.display(15.66, "minutes") == "15.7 min"
    assert f.display(None, "usd") == f.MISSING
    assert f.display(float("nan"), "trips") == f.MISSING


@pytest.mark.parametrize(
    ("ratio", "text"),
    [
        (0.5, "50%"),
        (0.1234, "12.34%"),
        (0.00123, "0.123%"),
        (0.00001, "<0.01%"),
        (0.99996, "99.996%"),
        (1, "100%"),
    ],
)
def test_percent_never_rounds_to_zero_or_a_hundred(ratio: float, text: str) -> None:
    assert f.percent(ratio) == text


def test_signed_change_and_spans() -> None:
    assert f.signed_change(0.045454) == "+4.5%"
    assert f.signed_change(-0.123) == "−12%"
    assert f.signed_change(0) == "0.0%"
    assert f.span_words(date(2025, 3, 1), date(2025, 3, 31)) == "from Mar 1 to Mar 31, 2025"
    assert f.date_range(date(2024, 12, 1), date(2025, 1, 31)) == "Dec 1, 2024 – Jan 31, 2025"
    assert f.plural(1, "file") == "1 file" and f.plural(6, "file") == "6 files"


def test_quality_labels_cover_every_pipeline_rule() -> None:
    assert set(QUARANTINE_LABELS) == set(QUARANTINE_REASONS)
    assert set(FLAG_LABELS) == set(FLAGS)


# ---- document --------------------------------------------------------------------------------------------


@pytest.mark.parametrize("template", list(TEMPLATES))
def test_every_template_builds_the_full_frame(template: str) -> None:
    doc = build(template)
    assert doc.title.endswith("test") and doc.period.label == "Mar 1 – Mar 31, 2025"
    assert doc.dataset.periods[0].run_id == str(RUN_ID) and len(doc.dataset.version_id) == 12
    assert doc.summary and doc.findings and doc.methodology and doc.limitations
    assert [s.id for s in doc.sections] == TEMPLATES[template].section_ids()
    assert doc.prepared_by == "Test Analyst"
    for finding in doc.findings:
        assert finding.evidence.source and finding.evidence.values  # every finding cites its values
    ReportDocument.model_validate_json(doc.model_dump_json())  # round-trips as the preview payload


def test_kpis_and_comparison_come_from_the_overview() -> None:
    doc = build("executive_overview")
    trips = next(k for k in doc.kpis if k.id == "total_trips")
    assert trips.value == TOTAL and trips.previous == TOTAL - 4000
    assert trips.change == pytest.approx(4000 / (TOTAL - 4000))
    assert trips.display == f"{TOTAL:,}" and trips.change_display == f.signed_change(trips.change)
    amount = next(k for k in doc.kpis if k.id == "avg_total_amount")
    assert amount.excluded_rows == 120 and amount.note == "Excludes 120 flagged trips"
    assert doc.comparison.available and doc.comparison.label == "Previous 31 days (Jan 29 – Feb 28, 2025)"
    first = doc.findings[0]
    assert first.section == "kpis" and first.evidence.values["previous_trips"] == TOTAL - 4000
    assert (
        doc.summary[0]
        == f"{TOTAL:,} trips were recorded from Mar 1 to Mar 31, 2025, {TOTAL / 31:,.0f} per day on average."
    )


def test_no_comparison_without_an_explicit_date_range() -> None:
    doc = build("executive_overview", AnalyticsFilters())
    assert not doc.comparison.available and "start and end date" in (doc.comparison.reason or "")
    assert all(k.previous is None for k in doc.kpis)
    assert doc.filters[1].value.startswith("All published data")


def test_findings_are_computed_from_the_returned_values() -> None:
    doc = build("demand_patterns")
    busiest = max(MARCH, key=_trips)
    trend = next(x for x in doc.findings if x.section == "trend")
    assert trend.evidence.values["busiest_date"] == busiest.isoformat()
    assert f"{_trips(busiest):,} trips" in trend.statement
    weekday = next(x for x in doc.findings if x.section == "weekday")
    per_day = {
        wd: sum(_trips(d) for d in MARCH if d.isoweekday() == wd)
        / sum(1 for d in MARCH if d.isoweekday() == wd)
        for wd in range(1, 8)
    }
    best = max(per_day, key=lambda wd: per_day[wd])
    assert weekday.evidence.values["busiest_per_day"] == pytest.approx(per_day[best], abs=0.1)
    assert weekday.statement.startswith(
        ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday")[best - 1]
    )
    peak = next(x for x in doc.findings if x.section == "hour")
    assert peak.evidence.values["peak_hour"] == 23 and "23:00 hour" in peak.statement
    slot = next(x for x in doc.findings if x.section == "heatmap")
    assert slot.evidence.values == {"weekday": "Sunday", "hour": 23, "trips": 723}


def test_sections_can_be_switched_off() -> None:
    spec = TEMPLATES["executive_overview"]
    doc = ReportBuilder(FakeAnalytics()).build(  # type: ignore[arg-type]
        template=spec.id,
        title="Short",
        filters=AnalyticsFilters(),
        sections=["weekday"],
        prepared_by="x",
    )
    assert [s.id for s in doc.sections] == ["weekday"]
    assert {x.section for x in doc.findings} == {"weekday"}


def test_quality_report_counts_only_the_months_it_covers() -> None:
    doc = build("data_quality")
    rows_read = next(k for k in doc.kpis if k.id == "rows_read")
    assert rows_read.value == TOTAL + 12  # the April run (99,999 rows) is outside the period
    assert not doc.comparison.available and doc.comparison.reason is None
    quarantine = next(s for s in doc.sections if s.id == "quarantine").blocks[0]
    assert isinstance(quarantine, TableBlock)
    assert quarantine.rows[0][:2] == ["Drop-off before pickup", 9]


def test_unknown_template_and_empty_results_are_refused() -> None:
    with pytest.raises(ValidationFailedError, match="unknown report template"):
        ReportBuilder(FakeAnalytics()).build(  # type: ignore[arg-type]
            template="nope", title="x", filters=AnalyticsFilters(), sections=[], prepared_by="x"
        )
    with pytest.raises(ValidationFailedError, match="no trips match"):
        build("executive_overview", empty=True)


# ---- renderers -------------------------------------------------------------------------------------------


@pytest.mark.parametrize("template", list(TEMPLATES))
def test_pdf_is_readable_with_page_numbers_attribution_and_methodology(template: str) -> None:
    doc = build(template)
    pdf, pages = render_pdf(doc)
    reader = PdfReader(io.BytesIO(pdf))
    assert len(reader.pages) == pages >= 2
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert doc.title in text and f"Page 1 of {pages}" in text and f"Page {pages} of {pages}" in text
    assert "Methodology" in text and "Limitations" in text and "Dataset version" in text
    assert "NYC TLC Trip Record Data (test)" in text
    assert reader.metadata is not None and reader.metadata.title == doc.title
    fonts = {
        str(font["/BaseFont"])
        for page in reader.pages
        for font in (page["/Resources"].get("/Font") or {}).values()  # type: ignore[union-attr]
        for font in [font.get_object()]
    }
    assert any("Geist" in name for name in fonts)  # embedded, not a substituted system font


def test_pdf_escapes_markup_from_data() -> None:
    pdf, _ = render_pdf(build("zone_analysis"))
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert MARKUP_ZONE in text  # shown literally, not interpreted as ReportLab markup


def test_xlsx_has_typed_cells_frozen_headers_and_no_formulas() -> None:
    doc = build("zone_analysis")
    book = openpyxl.load_workbook(io.BytesIO(render_report_xlsx(doc)))
    assert book.sheetnames[0] == "Summary" and {"Findings", "Dataset version", "Methodology"} <= set(
        book.sheetnames
    )
    summary = book["Summary"]
    trips_row = next(r for r in summary.iter_rows() if r[0].value == "Total trips")
    assert trips_row[1].value == TOTAL and trips_row[1].number_format == "#,##0"
    assert trips_row[3].number_format == "0.00%"
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                assert c.data_type != "f", f"{sheet.title}!{c.coordinate} became a formula: {c.value}"
    pairs = next(s for s in book.worksheets if s.title.startswith("Top 2 pickup"))
    assert pairs.freeze_panes == "A2"
    assert any(c.value == EVIL_ZONE and c.data_type == "s" for row in pairs.iter_rows() for c in row)
    chart_sheets = [s for s in book.worksheets if s._charts]  # openpyxl exposes embedded charts only here
    assert chart_sheets, "native charts are embedded"


def test_csv_is_tidy_and_neutralises_formulas() -> None:
    doc = build("zone_analysis")
    rows = list(csv.DictReader(io.StringIO(render_report_csv(doc).decode())))
    assert list(rows[0]) == COLUMNS
    assert {"report", "summary", "findings"} <= {r["section"] for r in rows}
    kpi = next(
        r for r in rows if r["block"] == "kpis" and r["label"] == "Total trips" and r["field"] == "value"
    )
    assert int(kpi["value"]) == TOTAL and kpi["unit"] == "trips"
    texts = [r["value"] for r in rows] + [r["label"] for r in rows]
    assert f"'{EVIL_ZONE}" in texts and EVIL_ZONE not in texts


def test_sheet_names_are_valid_and_unique() -> None:
    used: set[str] = set()
    names = [
        sheet_name(t, used)
        for t in ["A/B:C", "x" * 40, "Summary", "summary", "Trips by weekday and pickup hour"]
    ]
    assert names[0] == "A B C" and len(names[1]) <= 31 and names[3] == "summary (2)"
    assert names[4] == "Trips by weekday and pickup…"


def test_trip_extract_xlsx_streams_typed_rows(tmp_path: Any) -> None:
    path = tmp_path / "extract.xlsx"
    rows = [
        (datetime(2025, 3, 1, 8, 15, tzinfo=UTC), 132, 3.5, 21.75, EVIL_ZONE, ["invalid_amount"]),
        (datetime(2025, 3, 1, 9, 0, tzinfo=UTC), 1, None, -5.0, "plain", []),
    ]
    columns = [
        "pickup_datetime",
        "pickup_location_id",
        "trip_distance",
        "total_amount",
        "note",
        "quality_flags",
    ]
    written = write_extract_xlsx(path, columns, iter(rows), about=[("Filters", "start_date=2025-03-01")])
    assert written == 2
    book = openpyxl.load_workbook(path)
    trips = book["Trips"]
    assert trips.freeze_panes == "A2" and [c.value for c in trips[1]] == columns
    assert trips["A2"].value == datetime(2025, 3, 1, 8, 15) and trips["D3"].value == -5.0
    assert trips["D2"].number_format == '"$"#,##0.00'
    assert trips["E2"].value == EVIL_ZONE and trips["E2"].data_type == "s"
    assert trips["F2"].value == "invalid_amount" and trips["C3"].value is None
    assert book["About"]["B3"].value == "start_date=2025-03-01"


def test_chart_blocks_survive_every_renderer() -> None:
    doc = build("data_quality")
    daily = next(
        b for s in doc.sections for b in s.blocks if isinstance(b, ChartBlock) and b.id == "daily-flags"
    )
    assert [s.name for s in daily.series] == ["No passenger count", "Implausible amount"]
    assert daily.series[0].values[0] == pytest.approx(30 / _trips(MARCH[0]))
    render_pdf(doc)
    render_report_xlsx(doc)
    render_report_csv(doc)


def test_api_template_ids_match_the_registry() -> None:
    import typing

    from tripscope.api.routers.reports import TemplateId

    assert set(typing.get_args(TemplateId)) == set(TEMPLATES)
