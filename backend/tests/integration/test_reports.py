"""Phase 4 against the real services: report definitions, permissions, previews that equal the analytics API,
background generation of PDF/XLSX/CSV through the report worker, authorized and audited downloads, stale runs,
and XLSX trip extracts (FR-09, spec §12.2)."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import uuid
from datetime import timedelta
from typing import Any

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pypdf import PdfReader
from sqlalchemy import text

from tests.integration.conftest import PASSWORD
from tripscope.api.app import create_app
from tripscope.core.passwords import hash_password
from tripscope.metadata.models import JobStatus, ReportRun, Role, User
from tripscope.pipeline.runner import run_source
from tripscope.reports.worker import ReportWorker

pytestmark = [pytest.mark.integration, pytest.mark.spark]

JANUARY = {"start_date": "2025-01-01", "end_date": "2025-01-31"}


@pytest.fixture(scope="module")
def app(env: dict[str, Any]) -> FastAPI:
    return create_app(env["settings"])


@pytest.fixture(scope="module")
def published(env: dict[str, Any]) -> Any:
    summary = run_source(env["deps"], "fixture-2025-01")
    with env["factory"]() as session, session.begin():  # a second analyst, to test ownership
        session.add(
            User(
                email="analyst2@test.local",
                display_name="Second Analyst",
                password_hash=hash_password(PASSWORD),
                role=Role.ANALYST,
            )
        )
    return summary


@pytest.fixture(scope="module")
def as_(app: FastAPI, published: Any) -> dict[str, TestClient]:
    clients = {}
    for name in ("admin", "analyst", "analyst2", "viewer"):
        client = TestClient(app)
        response = client.post(
            "/api/v1/auth/login", json={"email": f"{name}@test.local", "password": PASSWORD}
        )
        assert response.status_code == 200
        clients[name] = client
    return clients


@pytest.fixture(scope="module")
def worker(env: dict[str, Any], app: FastAPI) -> ReportWorker:
    return ReportWorker(
        session_factory=env["factory"],
        analytics=app.state.analytics,
        store_factory=lambda: env["store"],
        poll_seconds=0.1,
        heartbeat_seconds=0.2,
    )


def create(client: TestClient, **body: Any) -> dict[str, Any]:
    response = client.post(
        "/api/v1/reports", json={"template": "executive_overview", "filters": JANUARY} | body
    )
    assert response.status_code == 201, response.text
    return dict(response.json())


def audit_count(env: dict[str, Any], action: str, target: str) -> int:
    with env["factory"]() as session:
        return int(
            session.execute(
                text("SELECT count(*) FROM audit_events WHERE action = :a AND target_id = :t"),
                {"a": action, "t": target},
            ).scalar_one()
        )


def test_templates_and_validation(as_: dict[str, TestClient]) -> None:
    catalogue = as_["viewer"].get("/api/v1/reports/templates").json()
    assert [t["number"] for t in catalogue["templates"]] == [1, 2, 3, 4, 5]
    assert catalogue["formats"] == ["pdf", "xlsx", "csv"]
    analyst = as_["analyst"]
    assert as_["viewer"].post("/api/v1/reports", json={"template": "executive_overview"}).status_code == 403
    assert analyst.post("/api/v1/reports", json={"template": "nope"}).status_code == 422
    assert (
        analyst.post("/api/v1/reports", json={"template": "data_quality", "sections": ["trend"]}).status_code
        == 422
    )
    assert analyst.post("/api/v1/reports", json={"template": "data_quality", "sql": "1"}).status_code == 422
    bad_filter = {"template": "executive_overview", "filters": {"pickup_zone": [999]}}
    assert analyst.post("/api/v1/reports", json=bad_filter).status_code == 422
    report = create(analyst, sections=["weekday", "trend"])
    assert report["sections"] == ["trend", "weekday"]  # template order
    assert report["title"] == "Executive overview · Jan 1 – Jan 31, 2025"
    assert report["visibility"] == "private" and report["permissions"] == {"edit": True, "generate": True}


def test_preview_equals_the_analytics_api(
    as_: dict[str, TestClient], env: dict[str, Any], published: Any
) -> None:
    analyst = as_["analyst"]
    report = create(analyst, title="January check")
    doc = analyst.get(f"/api/v1/reports/{report['report_id']}/preview").json()
    overview = analyst.get("/api/v1/analytics/overview", params=JANUARY).json()
    kpis = {k["id"]: k for k in doc["kpis"]}
    for metric, kpi in overview["kpis"].items():
        assert kpis[metric]["value"] == pytest.approx(kpi["value"]), metric
        assert kpis[metric]["excluded_rows"] == kpi["excluded_rows"]
    assert doc["title"] == "January check" and doc["prepared_by"] == "analyst"
    assert [p["run_id"] for p in doc["dataset"]["periods"]] == [str(published.run_id)]
    # The previous 31 days start before published coverage, so the report says so instead of guessing.
    assert (
        doc["comparison"]["available"] is False and "before published coverage" in doc["comparison"]["reason"]
    )
    weekday = next(s for s in doc["sections"] if s["id"] == "weekday")
    table = next(b for b in weekday["blocks"] if b["type"] == "table")
    by_weekday = analyst.get("/api/v1/analytics/trips-by-weekday", params=JANUARY).json()["groups"]
    assert (
        sum(r[1] for r in table["rows"])
        == sum(g["trips"] for g in by_weekday)
        == kpis["total_trips"]["value"]
    )


def test_generate_every_format_download_and_audit(
    as_: dict[str, TestClient], env: dict[str, Any], worker: ReportWorker
) -> None:
    analyst = as_["analyst"]
    report = create(analyst, title="Files", visibility="shared")
    rid = report["report_id"]
    runs = {}
    for fmt in ("pdf", "xlsx", "csv"):
        response = analyst.post(f"/api/v1/reports/{rid}/generate", json={"format": fmt})
        assert response.status_code == 202 and response.json()["status"] == "queued"
        runs[fmt] = response.json()["run_id"]
    duplicate = analyst.post(f"/api/v1/reports/{rid}/generate", json={"format": "pdf"})
    assert duplicate.status_code == 409 and duplicate.json()["error"]["details"]["run_id"] == runs["pdf"]
    assert analyst.get(f"/api/v1/reports/{rid}/download?format=pdf").status_code == 404  # nothing ready yet
    assert analyst.get(f"/api/v1/reports/{rid}/download?run_id={runs['pdf']}").status_code == 409

    for _ in range(3):
        assert worker.run_once() is not None
    assert worker.run_once() is None
    detail = analyst.get(f"/api/v1/reports/{rid}").json()
    by_format = {r["format"]: r for r in detail["runs"]}
    assert {r["status"] for r in by_format.values()} == {"completed"}
    assert by_format["pdf"]["page_count"] >= 2 and by_format["pdf"]["dataset_version_id"]

    total = analyst.get("/api/v1/analytics/overview", params=JANUARY).json()["kpis"]["total_trips"]["value"]
    pdf = analyst.get(f"/api/v1/reports/{rid}/download?format=pdf")
    assert pdf.status_code == 200 and pdf.headers["content-type"] == "application/pdf"
    assert (
        pdf.headers["x-report-sha256"]
        == hashlib.sha256(pdf.content).hexdigest()
        == by_format["pdf"]["sha256"]
    )
    assert (
        'filename="tripscope-executive-overview-20250101-20250131.pdf"' in pdf.headers["content-disposition"]
    )
    text_ = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf.content)).pages)
    assert "Files" in text_ and f"{total:,}" in text_ and "Page 1 of" in text_

    xlsx = as_["viewer"].get(f"/api/v1/reports/{rid}/download?format=xlsx")  # shared: viewers may download
    assert xlsx.status_code == 200
    summary = openpyxl.load_workbook(io.BytesIO(xlsx.content))["Summary"]
    assert next(r for r in summary.iter_rows(values_only=True) if r[0] == "Total trips")[1] == total

    rows = list(csv.DictReader(io.StringIO(analyst.get(f"/api/v1/reports/{rid}/download?format=csv").text)))
    kpi = next(
        r for r in rows if r["block"] == "kpis" and r["label"] == "Total trips" and r["field"] == "value"
    )
    assert int(kpi["value"]) == total

    keys = env["store"].list_keys(f"reports/{rid}/{runs['pdf']}/")
    assert sorted(k.rsplit("/", 1)[1] for k in keys) == ["document.json", by_format["pdf"]["file_name"]]
    stored = json.loads(env["store"].get_bytes(f"reports/{rid}/{runs['pdf']}/document.json"))
    assert stored["kpis"][0]["value"] == total and stored["title"] == "Files"
    assert audit_count(env, "report.downloaded", rid) == 3
    assert audit_count(env, "report.generate_requested", rid) == 3


def test_permission_matrix(as_: dict[str, TestClient]) -> None:
    private = create(as_["analyst"], title="Private")["report_id"]
    shared = create(as_["analyst"], title="Shared", visibility="shared")["report_id"]
    viewer, other, admin = as_["viewer"], as_["analyst2"], as_["admin"]

    # A private report is invisible to everyone but its owner and administrators (404, not 403).
    for client in (viewer, other):
        assert client.get(f"/api/v1/reports/{private}").status_code == 404
        assert client.get(f"/api/v1/reports/{private}/preview").status_code == 404
        assert client.get(f"/api/v1/reports/{private}/download").status_code == 404
    assert other.post(f"/api/v1/reports/{private}/generate", json={"format": "pdf"}).status_code == 404
    assert admin.get(f"/api/v1/reports/{private}").json()["permissions"] == {"edit": True, "generate": True}
    listed = {r["report_id"] for r in viewer.get("/api/v1/reports").json()["reports"]}
    assert shared in listed and private not in listed
    assert private in {r["report_id"] for r in admin.get("/api/v1/reports").json()["reports"]}

    # Shared: viewers read; other analysts generate but cannot edit or delete; viewers cannot generate.
    assert viewer.get(f"/api/v1/reports/{shared}").json()["permissions"] == {"edit": False, "generate": False}
    assert viewer.post(f"/api/v1/reports/{shared}/generate", json={"format": "csv"}).status_code == 403
    assert viewer.patch(f"/api/v1/reports/{shared}", json={"title": "x"}).status_code == 403
    assert other.patch(f"/api/v1/reports/{shared}", json={"title": "x"}).status_code == 403
    assert other.delete(f"/api/v1/reports/{shared}").status_code == 403
    assert other.post(f"/api/v1/reports/{shared}/generate", json={"format": "csv"}).status_code == 202
    assert admin.patch(f"/api/v1/reports/{shared}", json={"title": "Renamed by admin"}).status_code == 200
    assert viewer.get("/api/v1/reports?scope=mine").json()["reports"] == []


def test_runs_snapshot_the_definition(
    as_: dict[str, TestClient], env: dict[str, Any], worker: ReportWorker
) -> None:
    analyst = as_["analyst"]
    rid = create(analyst, title="Before edit", template="demand_patterns")["report_id"]
    run_id = analyst.post(f"/api/v1/reports/{rid}/generate", json={"format": "csv"}).json()["run_id"]
    edited = analyst.patch(
        f"/api/v1/reports/{rid}", json={"title": "After edit", "sections": ["hour"]}
    ).json()
    assert edited["title"] == "After edit" and edited["sections"] == ["hour"]
    while worker.run_once() is not None:
        pass
    stored = json.loads(env["store"].get_bytes(f"reports/{rid}/{run_id}/document.json"))
    assert stored["title"] == "Before edit" and len(stored["sections"]) == 5
    assert audit_count(env, "report.updated", rid) == 1


def test_failures_are_reported_and_stale_runs_fail(
    as_: dict[str, TestClient], env: dict[str, Any], worker: ReportWorker
) -> None:
    analyst = as_["analyst"]
    empty = create(analyst, filters={"start_date": "2025-01-20", "end_date": "2025-01-25"})["report_id"]
    assert analyst.get(f"/api/v1/reports/{empty}/preview").status_code == 422
    analyst.post(f"/api/v1/reports/{empty}/generate", json={"format": "pdf"})
    while worker.run_once() is not None:
        pass
    run = analyst.get(f"/api/v1/reports/{empty}").json()["runs"][0]
    assert run["status"] == "failed" and "no trips match" in run["error_summary"]

    rid = create(analyst)["report_id"]
    stale_id = uuid.UUID(
        analyst.post(f"/api/v1/reports/{rid}/generate", json={"format": "pdf"}).json()["run_id"]
    )
    with env["factory"]() as session, session.begin():
        run_row = session.get(ReportRun, stale_id)
        assert run_row is not None
        run_row.status, run_row.worker_id = JobStatus.RUNNING, "gone:1"
        run_row.heartbeat_at = run_row.created_at - timedelta(minutes=10)
    worker.run_once()
    stale = next(
        r for r in analyst.get(f"/api/v1/reports/{rid}").json()["runs"] if r["run_id"] == str(stale_id)
    )
    assert stale["status"] == "failed" and "stopped responding" in stale["error_summary"]
    # A failed format can be generated again.
    assert analyst.post(f"/api/v1/reports/{rid}/generate", json={"format": "pdf"}).status_code == 202


def test_delete_removes_files_and_is_audited(
    as_: dict[str, TestClient], env: dict[str, Any], worker: ReportWorker
) -> None:
    analyst = as_["analyst"]
    rid = create(analyst, template="data_quality", filters={})["report_id"]
    analyst.post(f"/api/v1/reports/{rid}/generate", json={"format": "xlsx"})
    while worker.run_once() is not None:
        pass
    assert env["store"].list_keys(f"reports/{rid}/")
    assert analyst.delete(f"/api/v1/reports/{rid}").status_code == 204
    assert analyst.get(f"/api/v1/reports/{rid}").status_code == 404
    assert env["store"].list_keys(f"reports/{rid}/") == []
    assert audit_count(env, "report.deleted", rid) == 1


def test_xlsx_trip_extract_is_typed_and_audited(as_: dict[str, TestClient], env: dict[str, Any]) -> None:
    viewer = as_["viewer"]
    request = {
        "format": "xlsx",
        "filters": {"payment_type": [1]},
        "scope": {"sort": "total_amount", "order": "desc"},
        "columns": ["pickup_datetime", "pickup_location_id", "total_amount", "quality_flags"],
        "max_rows": 3,
    }
    response = viewer.post("/api/v1/exports", json=request)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert response.headers["x-exported-rows"] == "3" and response.headers["x-truncated"] == "true"
    book = openpyxl.load_workbook(io.BytesIO(response.content))
    trips = book["Trips"]
    assert [c.value for c in trips[1]] == request["columns"] and trips.freeze_panes == "A2"
    amounts = [trips.cell(row=r, column=3).value for r in range(2, 5)]
    assert all(isinstance(a, int | float) for a in amounts) and amounts == sorted(amounts, reverse=True)
    assert trips.cell(row=2, column=3).number_format == '"$"#,##0.00'
    about = {r[0]: r[1] for r in book["About"].iter_rows(min_row=3, values_only=True)}
    assert about["Rows in this file"] == "3" and about["Source"] == "fixture"
    with env["factory"]() as session:
        details = session.execute(
            text(
                "SELECT details FROM audit_events WHERE action = 'export.xlsx' "
                "ORDER BY occurred_at DESC LIMIT 1"
            )
        ).scalar_one()
    assert details["exported_rows"] == 3
